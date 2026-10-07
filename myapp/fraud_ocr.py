import os
import re
import io
import zipfile

import pandas as pd
import joblib
import pytesseract

from PIL import Image, ImageOps, ImageEnhance

from docx import Document
import fitz

from django.conf import settings


# =========================================================
# TESSERACT CONFIGURATION
# =========================================================

TESSERACT_PATH = r"C:\Program Files\Tesseract-OCR\tesseract.exe"

if os.path.exists(TESSERACT_PATH):
    pytesseract.pytesseract.tesseract_cmd = TESSERACT_PATH


# =========================================================
# MODEL
# =========================================================

MODEL_PATH = os.path.join(
    settings.BASE_DIR,
    "myapp",
    "fraud_detection_model.pkl"
)

model_pipeline = joblib.load(MODEL_PATH)


# =========================================================
# MODEL FEATURES
# =========================================================

FEATURES = [
    'months_as_customer',
    'age',
    'policy_number',
    'policy_bind_date',
    'policy_state',
    'policy_csl',
    'policy_deductable',
    'policy_annual_premium',
    'umbrella_limit',
    'insured_zip',
    'insured_sex',
    'insured_education_level',
    'insured_occupation',
    'insured_hobbies',
    'insured_relationship',
    'capital-gains',
    'capital-loss',
    'incident_date',
    'incident_type',
    'collision_type',
    'incident_severity',
    'authorities_contacted',
    'incident_state',
    'incident_city',
    'incident_location',
    'incident_hour_of_the_day',
    'number_of_vehicles_involved',
    'property_damage',
    'bodily_injuries',
    'witnesses',
    'police_report_available',
    'total_claim_amount',
    'injury_claim',
    'property_claim',
    'vehicle_claim',
    'auto_make',
    'auto_model',
    'auto_year'
]


DATE_COLS = [
    'policy_bind_date',
    'incident_date'
]


NUM_COLS = [
    'months_as_customer',
    'age',
    'policy_number',
    'policy_deductable',
    'policy_annual_premium',
    'umbrella_limit',
    'insured_zip',
    'capital-gains',
    'capital-loss',
    'incident_hour_of_the_day',
    'number_of_vehicles_involved',
    'bodily_injuries',
    'witnesses',
    'total_claim_amount',
    'injury_claim',
    'property_claim',
    'vehicle_claim',
    'auto_year'
]


# =========================================================
# IMAGE PREPROCESSING
# =========================================================

def preprocess_image(image):

    try:

        image = image.convert("RGB")

        # grayscale
        image = ImageOps.grayscale(image)

        # increase contrast
        image = ImageEnhance.Contrast(image).enhance(2.0)

        # resize for better OCR
        width, height = image.size

        if width < 1800:

            scale = 1800 / width

            image = image.resize(
                (
                    int(width * scale),
                    int(height * scale)
                )
            )

        return image

    except Exception:

        return image


# =========================================================
# OCR IMAGE
# =========================================================

def extract_text_from_image(image):

    try:

        image = preprocess_image(image)

        text = pytesseract.image_to_string(
            image,
            config="--psm 6"
        )

        return text.strip()

    except Exception as e:

        return f"OCR ERROR: {str(e)}"


# =========================================================
# IMAGE FILE
# =========================================================

def extract_text_from_image_file(file_path):

    image = Image.open(file_path)

    return extract_text_from_image(image)


# =========================================================
# PDF TEXT
# =========================================================

def extract_text_from_pdf(file_path):

    all_text = []

    try:

        pdf = fitz.open(file_path)

        for page in pdf:

            text = page.get_text()

            if text:

                all_text.append(text)

        pdf.close()

    except Exception:

        pass

    return "\n".join(all_text).strip()


# =========================================================
# PDF OCR
# =========================================================

def extract_text_from_scanned_pdf(file_path):

    all_text = []

    pdf = fitz.open(file_path)

    for page_number in range(len(pdf)):

        page = pdf[page_number]

        # Render page to image
        pix = page.get_pixmap(
            matrix=fitz.Matrix(2, 2),
            alpha=False
        )

        image_bytes = pix.tobytes("png")

        image = Image.open(
            io.BytesIO(image_bytes)
        )

        text = extract_text_from_image(image)

        if text:

            all_text.append(
                f"\n--- PDF PAGE {page_number + 1} ---\n"
            )

            all_text.append(text)

    pdf.close()

    return "\n".join(all_text).strip()


# =========================================================
# PDF PROCESSING
# =========================================================

def process_pdf(file_path):

    text = extract_text_from_pdf(file_path)

    # If PDF has selectable text
    if len(text.strip()) >= 30:

        return text

    # Otherwise scanned PDF
    return extract_text_from_scanned_pdf(
        file_path
    )


# =========================================================
# DOCX TEXT
# =========================================================

def extract_docx_text(file_path):

    all_text = []

    document = Document(file_path)

    # -----------------------------------------------------
    # Paragraphs
    # -----------------------------------------------------

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:

            all_text.append(text)

    # -----------------------------------------------------
    # Tables
    # -----------------------------------------------------

    for table in document.tables:

        for row in table.rows:

            row_values = []

            for cell in row.cells:

                value = cell.text.strip()

                if value:

                    row_values.append(value)

            if row_values:

                all_text.append(
                    " : ".join(row_values)
                )

    return "\n".join(all_text).strip()


# =========================================================
# DOCX EMBEDDED IMAGES OCR
# =========================================================

def extract_docx_images_ocr(file_path):

    image_text = []

    try:

        with zipfile.ZipFile(
            file_path,
            "r"
        ) as docx_zip:

            for name in docx_zip.namelist():

                if not name.startswith(
                    "word/media/"
                ):

                    continue

                image_data = docx_zip.read(name)

                try:

                    image = Image.open(
                        io.BytesIO(image_data)
                    )

                    text = extract_text_from_image(
                        image
                    )

                    if text:

                        image_text.append(
                            f"\n--- DOCX IMAGE: {name} ---\n"
                        )

                        image_text.append(text)

                except Exception:

                    continue

    except Exception:

        pass

    return "\n".join(image_text).strip()


# =========================================================
# DOCX PROCESSING
# =========================================================

def process_docx(file_path):

    normal_text = extract_docx_text(
        file_path
    )

    image_text = extract_docx_images_ocr(
        file_path
    )

    final_text = ""

    if normal_text:

        final_text += normal_text

    if image_text:

        final_text += "\n\n"
        final_text += image_text

    return final_text.strip()


# =========================================================
# NORMALIZE OCR TEXT
# =========================================================

def normalize_text(text):

    text = text.replace(
        "\r",
        "\n"
    )

    text = re.sub(
        r'[ \t]+',
        ' ',
        text
    )

    return text


# =========================================================
# FIND VALUE AFTER LABEL
# =========================================================

def find_value(text, labels):

    if isinstance(labels, str):

        labels = [labels]

    for label in labels:

        pattern = (
            r'(?:^|\n|\r)'
            r'\s*'
            + re.escape(label)
            + r'\s*[:\-]\s*'
            r'([^\n\r]+)'
        )

        match = re.search(
            pattern,
            text,
            flags=re.IGNORECASE
        )

        if match:

            value = match.group(1).strip()

            return value

    return ""


# =========================================================
# FIND NUMBER
# =========================================================

def find_number(text, labels):

    value = find_value(
        text,
        labels
    )

    if not value:

        return ""

    match = re.search(
        r'-?\d+(?:\.\d+)?',
        value.replace(",", "")
    )

    if match:

        return match.group(0)

    return ""


# =========================================================
# FIND DATE
# =========================================================

def find_date(text, labels):

    value = find_value(
        text,
        labels
    )

    if not value:

        return ""

    patterns = [
        r'\d{4}-\d{2}-\d{2}',
        r'\d{2}/\d{2}/\d{4}',
        r'\d{2}-\d{2}-\d{4}',
        r'\d{1,2}/\d{1,2}/\d{4}',
    ]

    for pattern in patterns:

        match = re.search(
            pattern,
            value
        )

        if match:

            raw_date = match.group(0)

            parsed = pd.to_datetime(
                raw_date,
                errors="coerce"
            )

            if not pd.isna(parsed):

                return parsed.strftime(
                    "%Y-%m-%d"
                )

    return ""


# =========================================================
# EXTRACT INSURANCE DATA
# =========================================================

def extract_insurance_data(text):

    text = normalize_text(text)

    data = {}

    # -----------------------------------------------------
    # Customer
    # -----------------------------------------------------

    data['months_as_customer'] = find_number(
        text,
        [
            'months_as_customer',
            'Months as Customer',
            'Months As Customer'
        ]
    )

    data['age'] = find_number(
        text,
        ['age', 'Age']
    )

    data['policy_number'] = find_number(
        text,
        [
            'policy_number',
            'Policy Number'
        ]
    )

    # -----------------------------------------------------
    # Policy
    # -----------------------------------------------------

    data['policy_bind_date'] = find_date(
        text,
        [
            'policy_bind_date',
            'Policy Bind Date'
        ]
    )

    data['policy_state'] = find_value(
        text,
        [
            'policy_state',
            'Policy State'
        ]
    )

    data['policy_csl'] = find_value(
        text,
        [
            'policy_csl',
            'Policy CSL'
        ]
    )

    data['policy_deductable'] = find_number(
        text,
        [
            'policy_deductable',
            'Policy Deductable',
            'Policy Deductible'
        ]
    )

    data['policy_annual_premium'] = find_number(
        text,
        [
            'policy_annual_premium',
            'Policy Annual Premium'
        ]
    )

    data['umbrella_limit'] = find_number(
        text,
        [
            'umbrella_limit',
            'Umbrella Limit'
        ]
    )

    # -----------------------------------------------------
    # Insured
    # -----------------------------------------------------

    data['insured_zip'] = find_number(
        text,
        [
            'insured_zip',
            'Insured Zip'
        ]
    )

    data['insured_sex'] = find_value(
        text,
        [
            'insured_sex',
            'Insured Sex'
        ]
    )

    data['insured_education_level'] = find_value(
        text,
        [
            'insured_education_level',
            'Insured Education Level'
        ]
    )

    data['insured_occupation'] = find_value(
        text,
        [
            'insured_occupation',
            'Insured Occupation'
        ]
    )

    data['insured_hobbies'] = find_value(
        text,
        [
            'insured_hobbies',
            'Insured Hobbies'
        ]
    )

    data['insured_relationship'] = find_value(
        text,
        [
            'insured_relationship',
            'Insured Relationship'
        ]
    )

    # -----------------------------------------------------
    # Financial
    # -----------------------------------------------------

    data['capital-gains'] = find_number(
        text,
        [
            'capital-gains',
            'Capital Gains',
            'Capital-gains'
        ]
    )

    data['capital-loss'] = find_number(
        text,
        [
            'capital-loss',
            'Capital Loss',
            'Capital-loss'
        ]
    )

    # -----------------------------------------------------
    # Incident
    # -----------------------------------------------------

    data['incident_date'] = find_date(
        text,
        [
            'incident_date',
            'Incident Date'
        ]
    )

    data['incident_type'] = find_value(
        text,
        [
            'incident_type',
            'Incident Type'
        ]
    )

    data['collision_type'] = find_value(
        text,
        [
            'collision_type',
            'Collision Type'
        ]
    )

    data['incident_severity'] = find_value(
        text,
        [
            'incident_severity',
            'Incident Severity'
        ]
    )

    data['authorities_contacted'] = find_value(
        text,
        [
            'authorities_contacted',
            'Authorities Contacted'
        ]
    )

    data['incident_state'] = find_value(
        text,
        [
            'incident_state',
            'Incident State'
        ]
    )

    data['incident_city'] = find_value(
        text,
        [
            'incident_city',
            'Incident City'
        ]
    )

    data['incident_location'] = find_value(
        text,
        [
            'incident_location',
            'Incident Location'
        ]
    )

    data['incident_hour_of_the_day'] = find_number(
        text,
        [
            'incident_hour_of_the_day',
            'Incident Hour of the Day'
        ]
    )

    data['number_of_vehicles_involved'] = find_number(
        text,
        [
            'number_of_vehicles_involved',
            'Number of Vehicles Involved'
        ]
    )

    # -----------------------------------------------------
    # Damage / Injury
    # -----------------------------------------------------

    data['property_damage'] = find_value(
        text,
        [
            'property_damage',
            'Property Damage'
        ]
    )

    data['bodily_injuries'] = find_number(
        text,
        [
            'bodily_injuries',
            'Bodily Injuries'
        ]
    )

    data['witnesses'] = find_number(
        text,
        [
            'witnesses',
            'Witnesses'
        ]
    )

    data['police_report_available'] = find_value(
        text,
        [
            'police_report_available',
            'Police Report Available'
        ]
    )

    # -----------------------------------------------------
    # Claim
    # -----------------------------------------------------

    data['total_claim_amount'] = find_number(
        text,
        [
            'total_claim_amount',
            'Total Claim Amount'
        ]
    )

    data['injury_claim'] = find_number(
        text,
        [
            'injury_claim',
            'Injury Claim'
        ]
    )

    data['property_claim'] = find_number(
        text,
        [
            'property_claim',
            'Property Claim'
        ]
    )

    data['vehicle_claim'] = find_number(
        text,
        [
            'vehicle_claim',
            'Vehicle Claim'
        ]
    )

    # -----------------------------------------------------
    # Vehicle
    # -----------------------------------------------------

    data['auto_make'] = find_value(
        text,
        [
            'auto_make',
            'Auto Make'
        ]
    )

    data['auto_model'] = find_value(
        text,
        [
            'auto_model',
            'Auto Model'
        ]
    )

    data['auto_year'] = find_number(
        text,
        [
            'auto_year',
            'Auto Year'
        ]
    )

    # Make sure every feature exists
    for feature in FEATURES:

        if feature not in data:

            data[feature] = ""

    return data


# =========================================================
# PREPARE MODEL DATA
# =========================================================

def prepare_model_data(extracted_data):

    model_data = {}

    for feature in FEATURES:

        value = extracted_data.get(
            feature,
            ""
        )

        model_data[feature] = value

    # -----------------------------------------------------
    # Date conversion
    # -----------------------------------------------------

    for col in DATE_COLS:

        value = model_data.get(
            col,
            ""
        )

        if value:

            try:

                date_value = pd.to_datetime(
                    value,
                    errors="coerce"
                )

                if not pd.isna(date_value):

                    model_data[col] = (
                        date_value
                        - pd.Timestamp("1899-12-30")
                    ).days

                else:

                    model_data[col] = 0

            except Exception:

                model_data[col] = 0

        else:

            model_data[col] = 0

    # -----------------------------------------------------
    # Numeric conversion
    # -----------------------------------------------------

    for col in NUM_COLS:

        value = model_data.get(
            col,
            ""
        )

        if value == "":

            model_data[col] = 0

            continue

        try:

            clean_value = str(value).replace(
                ",",
                ""
            )

            model_data[col] = float(
                clean_value
            )

        except Exception:

            model_data[col] = 0

    return pd.DataFrame(
        [model_data],
        columns=FEATURES
    )


# =========================================================
# PREDICT FRAUD
# =========================================================

def predict_fraud(extracted_data):

    new_df = prepare_model_data(
        extracted_data
    )

    prediction = model_pipeline.predict(
        new_df
    )

    probability = None

    try:

        probability = (
            model_pipeline
            .predict_proba(new_df)[0][1]
        )

    except Exception:

        probability = None

    prediction_value = int(
        prediction[0]
    )

    if probability is not None:

        probability_percentage = round(
            float(probability) * 100,
            2
        )

    else:

        probability_percentage = None

    return {

        'prediction_value':
            prediction_value,

        'prediction':
            (
                'Fraud Detected'
                if prediction_value == 1
                else 'No Fraud'
            ),

        'probability':
            probability_percentage
    }


# =========================================================
# MAIN DOCUMENT PROCESSOR
# =========================================================

def process_uploaded_document(
    file_path,
    extension=None
):

    if extension is None:

        extension = os.path.splitext(
            file_path
        )[1].lower()

    # -----------------------------------------------------
    # IMAGE
    # -----------------------------------------------------

    if extension in [
        '.jpg',
        '.jpeg',
        '.png'
    ]:

        ocr_text = extract_text_from_image_file(
            file_path
        )

    # -----------------------------------------------------
    # PDF
    # -----------------------------------------------------

    elif extension == '.pdf':

        ocr_text = process_pdf(
            file_path
        )

    # -----------------------------------------------------
    # DOCX
    # -----------------------------------------------------

    elif extension == '.docx':

        ocr_text = process_docx(
            file_path
        )

    else:

        raise ValueError(
            "Unsupported file format."
        )

    if not ocr_text:

        raise ValueError(
            "No readable text could be extracted "
            "from the uploaded document."
        )

    # -----------------------------------------------------
    # Extract fields
    # -----------------------------------------------------

    extracted_data = extract_insurance_data(
        ocr_text
    )

    # -----------------------------------------------------
    # Model prediction
    # -----------------------------------------------------

    result = predict_fraud(
        extracted_data
    )

    return {

        'ocr_text':
            ocr_text,

        'extracted_data':
            extracted_data,

        'result':
            result
    }